#!/usr/bin/env python3
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from scipy.stats import fisher_exact, mannwhitneyu, ttest_ind

ROOT = Path(__file__).resolve().parent
INPUT_XLSX = ROOT / "20260826_DeID.xlsx"
OUTPUT_DOCX = ROOT / "20260919_Supplementary_Table1.docx"
OUTPUT_AUDIT_CSV = ROOT / "20260919_Supplementary_Table1_Cohort_Reconciliation.csv"
OUTPUT_METADATA_CSV = ROOT / "20260919_Supplementary_Table1_Cohort_Metadata.csv"

COMPLETION_COL = "PAC_Program_Completion"
COMPLETER_LABEL = "Completed PAC program"
ID_COL = "ID"
AMBULATORY_FLAG_COL = "6MWT_Best_Scenario"
AMBULATORY_DISTANCE_COL = "Initial_6MWT_Distance"
OUTCOME_COL = "6MWT4"

IN_NIHSS_COLS = [
    "ConsIn", "AnswerIn", "OrderIn", "EOMIn", "VisualIn", "FaceIn", "LUIn", "RUIn",
    "LLIn", "RLIn", "CoordinateIn", "SensoryIn", "LanguageIn", "ArticulateIn", "NeglectIn",
]
OUT_NIHSS_COLS = [
    "ConsOut", "AnswerOut", "OrderOut", "EOMOut", "VisualOut", "FacialOut", "LUOut", "RUOut",
    "LLOut", "RLOut", "Coordinateout", "SensoryOut", "LanguageOut", "ArticulateOut", "NeglectOut",
]


def fmt_p(value: float) -> str:
    if pd.isna(value):
        return "—"
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def fmt_mean_sd(series: pd.Series) -> str:
    x = pd.to_numeric(series, errors="coerce").dropna()
    if x.empty:
        return "—"
    return f"{x.mean():.1f} ± {x.std(ddof=0):.1f}"


def fmt_median_iqr(series: pd.Series) -> str:
    x = pd.to_numeric(series, errors="coerce").dropna()
    if x.empty:
        return "—"
    q1 = x.quantile(0.25)
    q3 = x.quantile(0.75)
    return f"{x.median():.1f} [{q1:.1f}, {q3:.1f}]"


def fmt_n_pct(binary_series: pd.Series) -> str:
    x = pd.to_numeric(binary_series, errors="coerce")
    denom = x.notna().sum()
    if denom == 0:
        return "—"
    num = int((x == 1).sum())
    return f"{num} ({100.0 * num / denom:.1f}%)"


def p_ttest(a: pd.Series, b: pd.Series) -> float:
    aa = pd.to_numeric(a, errors="coerce").dropna()
    bb = pd.to_numeric(b, errors="coerce").dropna()
    if aa.empty or bb.empty:
        return np.nan
    return float(ttest_ind(aa, bb, equal_var=False, nan_policy="omit").pvalue)


def p_mannwhitney(a: pd.Series, b: pd.Series) -> float:
    aa = pd.to_numeric(a, errors="coerce").dropna()
    bb = pd.to_numeric(b, errors="coerce").dropna()
    if aa.empty or bb.empty:
        return np.nan
    return float(mannwhitneyu(aa, bb, alternative="two-sided").pvalue)


def p_fisher(binary_a: pd.Series, binary_b: pd.Series) -> float:
    aa = pd.to_numeric(binary_a, errors="coerce")
    bb = pd.to_numeric(binary_b, errors="coerce")
    a1 = int((aa == 1).sum())
    a0 = int((aa == 0).sum())
    b1 = int((bb == 1).sum())
    b0 = int((bb == 0).sum())
    if (a1 + a0 == 0) or (b1 + b0 == 0):
        return np.nan
    return float(fisher_exact([[a1, a0], [b1, b0]], alternative="two-sided").pvalue)


def derive_cohort(df: pd.DataFrame) -> pd.DataFrame:
    work = df.copy()
    work["is_completer"] = work[COMPLETION_COL].eq(COMPLETER_LABEL)
    work["is_eligible_ambulatory"] = work["is_completer"] & pd.to_numeric(
        work[AMBULATORY_FLAG_COL], errors="coerce"
    ).eq(1)
    # Cross-check against the historical ambulatory-distance marker to avoid silent drift.
    ambulatory_distance = work["is_completer"] & work[AMBULATORY_DISTANCE_COL].notna()
    if not work["is_eligible_ambulatory"].equals(ambulatory_distance):
        raise ValueError(
            "Ambulatory eligibility mismatch between 6MWT_Best_Scenario==1 and non-missing Initial_6MWT_Distance."
        )
    work["has_observed_outcome"] = pd.to_numeric(work[OUTCOME_COL], errors="coerce").notna()
    work["in_model_analytic"] = work["is_eligible_ambulatory"] & work["has_observed_outcome"]
    work["excluded_from_model"] = work["is_eligible_ambulatory"] & ~work["has_observed_outcome"]
    return work


def build_metadata(work: pd.DataFrame) -> dict:
    n_eligible = int(work["is_eligible_ambulatory"].sum())
    n_analytic = int(work["in_model_analytic"].sum())
    n_excluded = int(work["excluded_from_model"].sum())
    if n_analytic + n_excluded != n_eligible:
        raise ValueError("Cohort reconciliation failed: analytic N + excluded N does not equal eligible N.")
    return {
        "Source workbook": INPUT_XLSX.name,
        "Eligible ambulatory definition": "Completed PAC program AND 6MWT_Best_Scenario == 1",
        "Model-analytic definition": "Eligible ambulatory AND non-missing 6MWT4",
        "N eligible ambulatory": n_eligible,
        "N model-analytic (observed 6MWT4)": n_analytic,
        "N excluded from model": n_excluded,
        "Excluded reason": "Missing 6MWT4",
        "Excluded IDs (de-identified)": ",".join(map(str, work.loc[work["excluded_from_model"], ID_COL].astype(int).tolist())),
    }


def build_exclusion_audit(work: pd.DataFrame) -> pd.DataFrame:
    audit = work.loc[
        work["excluded_from_model"],
        [ID_COL, COMPLETION_COL, AMBULATORY_FLAG_COL, AMBULATORY_DISTANCE_COL, OUTCOME_COL],
    ].copy()
    audit["Exclusion_Condition"] = "Eligible ambulatory completer but missing 6MWT4 (not in continuous model)"
    return audit.sort_values(ID_COL).reset_index(drop=True)


def main() -> None:
    df = pd.read_excel(INPUT_XLSX)
    work = derive_cohort(df)
    metadata = build_metadata(work)
    audit_df = build_exclusion_audit(work)
    pd.DataFrame([metadata]).to_csv(OUTPUT_METADATA_CSV, index=False)
    audit_df.to_csv(OUTPUT_AUDIT_CSV, index=False)

    # Derived binary variables
    work["Male"] = pd.to_numeric(work["Sex, F0 M1"], errors="coerce")
    work["IschemicStroke"] = np.where(pd.to_numeric(work["HemorrhageStroke"], errors="coerce").isna(), np.nan,
                                      (pd.to_numeric(work["HemorrhageStroke"], errors="coerce") == 0).astype(int))
    work["GIB_Present"] = np.where(pd.to_numeric(work["GIB"], errors="coerce").isna(), np.nan,
                                   (pd.to_numeric(work["GIB"], errors="coerce") > 0).astype(int))

    in_total = work[IN_NIHSS_COLS].apply(pd.to_numeric, errors="coerce").sum(axis=1, min_count=1)
    out_total = work[OUT_NIHSS_COLS].apply(pd.to_numeric, errors="coerce").sum(axis=1, min_count=1)

    analytic = work.loc[work["in_model_analytic"]].copy()
    excluded = work.loc[work["excluded_from_model"]].copy()

    rows = [
        (
            "Baseline Age (years, mean ± SD)",
            fmt_mean_sd(analytic["Age"]),
            fmt_mean_sd(excluded["Age"]),
            fmt_p(p_ttest(analytic["Age"], excluded["Age"])),
        ),
        (
            "Male Sex (n, %)",
            fmt_n_pct(analytic["Male"]),
            fmt_n_pct(excluded["Male"]),
            fmt_p(p_fisher(analytic["Male"], excluded["Male"])),
        ),
        (
            "Ischemic Stroke Etiology (n, %)",
            fmt_n_pct(analytic["IschemicStroke"]),
            fmt_n_pct(excluded["IschemicStroke"]),
            fmt_p(p_fisher(analytic["IschemicStroke"], excluded["IschemicStroke"])),
        ),
        (
            "Baseline Barthel Index (median, IQR)",
            fmt_median_iqr(analytic["BI1"]),
            fmt_median_iqr(excluded["BI1"]),
            fmt_p(p_mannwhitney(analytic["BI1"], excluded["BI1"])),
        ),
        (
            "Baseline Berg Balance Scale (mean ± SD)",
            fmt_mean_sd(analytic["BBS1"]),
            fmt_mean_sd(excluded["BBS1"]),
            fmt_p(p_ttest(analytic["BBS1"], excluded["BBS1"])),
        ),
        (
            "Acute Admission NIHSS (median, IQR)",
            fmt_median_iqr(in_total.loc[analytic.index]),
            fmt_median_iqr(in_total.loc[excluded.index]),
            fmt_p(p_mannwhitney(in_total.loc[analytic.index], in_total.loc[excluded.index])),
        ),
        (
            "Acute Discharge NIHSS (median, IQR)",
            fmt_median_iqr(out_total.loc[analytic.index]),
            fmt_median_iqr(out_total.loc[excluded.index]),
            fmt_p(p_mannwhitney(out_total.loc[analytic.index], out_total.loc[excluded.index])),
        ),
        (
            "Gastrointestinal Bleeding (GIB) (n, %)",
            fmt_n_pct(analytic["GIB_Present"]),
            fmt_n_pct(excluded["GIB_Present"]),
            fmt_p(p_fisher(analytic["GIB_Present"], excluded["GIB_Present"])),
        ),
    ]

    doc = Document()
    title = doc.add_paragraph(
        "Supplementary Table 1. Ambulatory-cohort reconciliation for the 6MWT4 continuous model"
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(11)

    n_eligible = metadata["N eligible ambulatory"]
    n_analytic = metadata["N model-analytic (observed 6MWT4)"]
    n_excluded = metadata["N excluded from model"]
    doc.add_paragraph(
        f"Source workbook: {INPUT_XLSX.name}. Eligible ambulatory cohort: n={n_eligible} "
        f"(Completed PAC program and 6MWT_Best_Scenario=1). "
        f"Continuous-model analytic cohort: n={n_analytic} (non-missing 6MWT4). "
        f"Excluded from model: n={n_excluded} (missing 6MWT4)."
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = [
        "Characteristic",
        f"Model-analytic ambulatory (n={n_analytic})",
        f"Eligible ambulatory excluded (n={n_excluded})",
        "P value",
    ]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)

    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9)

    doc.add_paragraph(
        "Continuous mean±SD variables were compared with Welch t-test; median [IQR] variables with Mann–Whitney U test; categorical variables with Fisher's exact test. "
        "Ischemic stroke etiology was derived as HemorrhageStroke=0. Admission/discharge NIHSS were computed as sums of NIHSS item scores (In/Out components). "
        f"Cohort reconciliation: {n_eligible} eligible ambulatory participants = {n_analytic} model-analytic + {n_excluded} excluded for missing 6MWT4. "
        f"Excluded IDs (de-identified): {metadata['Excluded IDs (de-identified)']}. "
        f"Record-level audit: {OUTPUT_AUDIT_CSV.name}; metadata: {OUTPUT_METADATA_CSV.name}."
    )

    doc.save(OUTPUT_DOCX)
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Saved: {OUTPUT_AUDIT_CSV}")
    print(f"Saved: {OUTPUT_METADATA_CSV}")


if __name__ == "__main__":
    main()
