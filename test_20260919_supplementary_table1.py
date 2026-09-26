import importlib.util
import unittest
from pathlib import Path

import pandas as pd
from docx import Document


MODULE_PATH = Path(__file__).resolve().parent / "20260919_Supplementary_Table1.py"
SPEC = importlib.util.spec_from_file_location("supp_table1_20260919", MODULE_PATH)
supp_table1 = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(supp_table1)


class SupplementaryTable1CohortTests(unittest.TestCase):
    def test_reconciliation_counts_and_ids(self):
        df = pd.read_excel(supp_table1.INPUT_XLSX)
        work = supp_table1.derive_cohort(df)
        metadata = supp_table1.build_metadata(work)
        audit = supp_table1.build_exclusion_audit(work)

        self.assertEqual(metadata["N eligible ambulatory"], 517)
        self.assertEqual(metadata["N model-analytic (observed 6MWT4)"], 511)
        self.assertEqual(metadata["N excluded from model"], 6)
        self.assertEqual(metadata["Excluded IDs (de-identified)"], "107,112,208,257,458,550")
        self.assertEqual(metadata["N model-analytic (observed 6MWT4)"] + metadata["N excluded from model"], metadata["N eligible ambulatory"])

        self.assertEqual(audit[supp_table1.ID_COL].astype(int).tolist(), [107, 112, 208, 257, 458, 550])
        self.assertTrue((pd.to_numeric(audit[supp_table1.OUTCOME_COL], errors="coerce").isna()).all())

    def test_generated_table_labels_match_metadata(self):
        supp_table1.main()
        metadata_df = pd.read_csv(supp_table1.OUTPUT_METADATA_CSV)
        metadata = metadata_df.iloc[0].to_dict()

        doc = Document(supp_table1.OUTPUT_DOCX)
        self.assertIn("Ambulatory-cohort reconciliation", doc.paragraphs[0].text)
        header_cells = [c.text for c in doc.tables[0].rows[0].cells]
        self.assertEqual(header_cells[1], f"Model-analytic ambulatory (n={int(metadata['N model-analytic (observed 6MWT4)'])})")
        self.assertEqual(header_cells[2], f"Eligible ambulatory excluded (n={int(metadata['N excluded from model'])})")


if __name__ == "__main__":
    unittest.main()
